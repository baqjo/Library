"""Exports (Excel / CSV) of statistics and operational lists."""
import csv
import io
from datetime import datetime
from flask import Blueprint, render_template, request, Response, abort
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from . import db, circ
from .auth import staff
from .i18n import tr
from .models import Loan, Hold, Fine, Copy, Title, Student, Enrollment, Semester
from .stats import compute, STAT_COLS
from .util import L, parse_date

bp = Blueprint("reports", __name__)


def _safe(v):
    """Neutralise spreadsheet formulas in text that came from users (names, titles, imports)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _range(args):
    return parse_date(args.get("from")), parse_date(args.get("to"))


def _in_range(dt, start, end):
    d = dt.date() if isinstance(dt, datetime) else dt
    return (not start or d >= start) and (not end or d <= end)


def _sec_label(sec, lang):
    return sec.label(lang) if sec else ""


# ---- each table: (title_key, header_keys, rows) ----
def t_stats(args):
    res = compute(args)
    dim = res["f"]["dim"]
    heads = ["dim_" + dim] + [c[0] for c in STAT_COLS]
    rows = [[r["label"]] + [r[c[1]] for c in STAT_COLS] for r in res["rows"]]
    tot = res["totals"]
    rows.append([tr("total", L())] + [tot["students"], tot["loans"], tot["borrowers"], tot["participation"], None,
                                     100.0 if tot["loans"] else None, None, tot["holds"], tot["fines"], tot["unpaid"]])
    return "stats", heads, rows


def t_overdue(args):
    lang, t = L(), circ.today()
    rows = []
    for l in db.session.query(Loan).filter(Loan.returned_at.is_(None), Loan.due_date < t).order_by(Loan.due_date):
        s = l.student
        rows.append([s.student_no, s.name(lang), _sec_label(l.section, lang), l.copy.title.title, l.copy.barcode,
                     l.due_date, (t - l.due_date).days, s.email])
    return "overdue_list", ["student_no", "name", "class_col", "title", "copy_barcode", "due", "days_overdue", "email"], rows


def t_loans(args):
    lang, (a, b) = L(), _range(args)
    rows = []
    for l in Loan.query.order_by(Loan.out_at.desc()).all():
        if not _in_range(l.out_at, a, b):
            continue
        s = l.student
        rows.append([s.student_no, s.name(lang), _sec_label(l.section, lang), l.copy.title.title, l.copy.barcode,
                     l.out_at.date(), l.due_date, l.returned_at.date() if l.returned_at else None, l.renewals])
    return "loans_list", ["student_no", "name", "class_col", "title", "copy_barcode", "out_date", "due", "returned", "renewals"], rows


def t_holds(args):
    lang, (a, b) = L(), _range(args)
    rows = []
    for h in Hold.query.order_by(Hold.id.desc()).all():
        if not _in_range(h.created_at, a, b):
            continue
        rows.append([h.student.student_no, h.student.name(lang), h.title.title, tr("hold_" + h.status if h.status in ("waiting", "ready", "fulfilled", "expired") else "hold_cancelled_s", lang),
                     h.created_at.date(), h.expires_at, h.copy.barcode if h.copy else ""])
    return "holds", ["student_no", "name", "title", "status", "date", "pickup_until", "copy_barcode"], rows


def t_fines(args):
    lang, (a, b) = L(), _range(args)
    rows = []
    for f in Fine.query.order_by(Fine.id.desc()).all():
        if not _in_range(f.created_at, a, b):
            continue
        rows.append([f.student.student_no, f.student.name(lang), f.loan.copy.title.title if f.loan else "",
                     tr("overdue_r" if f.reason == "overdue" else "lost_r", lang), f.amount, tr(f.status, lang), f.created_at.date()])
    return "fines", ["student_no", "name", "title", "reason", "amount", "status", "date"], rows


def t_inventory(args):
    rows = [[c.barcode, c.title.title, c.title.author, c.title.isbn or "", c.title.category, c.title.call_number,
             c.shelf, c.location, tr(c.status if c.status != "out" else "out", L()), c.price]
            for c in Copy.query.join(Title, Title.id == Copy.title_id).order_by(Copy.barcode).all()]
    return "inventory", ["copy_barcode", "title", "author", "isbn", "category", "call_number", "shelf", "location", "status", "price"], rows


def t_students(args):
    lang, y = L(), circ.current_year()
    en = {e.student_id: e for e in Enrollment.query.filter_by(year_id=y.id).all()} if y else {}
    rows = [[s.student_no, s.name_ar, s.name_en, s.email, s.barcode or "",
             _sec_label(en[s.id].section, lang) if s.id in en and en[s.id].section else ""]
            for s in Student.query.order_by(Student.student_no).all()]
    return "students", ["student_no", "name_ar", "name_en", "email", "barcode", "class_col"], rows


TABLES = {"stats": t_stats, "overdue": t_overdue, "loans": t_loans, "holds": t_holds, "fines": t_fines,
          "inventory": t_inventory, "students": t_students}


def to_xlsx(title_key, headers, rows, lang):
    wb = Workbook()
    ws = wb.active
    ws.title = tr(title_key, lang)[:31].replace("/", "-")
    ws.sheet_view.rightToLeft = lang == "ar"
    ws.append([tr(h, lang) for h in headers])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_safe(v) for v in r])
    for i in range(1, len(headers) + 1):
        width = max([len(str(ws.cell(row=r, column=i).value or "")) for r in range(1, min(ws.max_row, 200) + 1)] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv(headers, rows, lang):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([tr(h, lang) for h in headers])
    for r in rows:
        w.writerow(["" if v is None else _safe(v) for v in r])
    return ("﻿" + out.getvalue()).encode("utf-8")


@bp.route("/reports")
@staff
def index():
    return render_template("reports.html", tables=[k for k in TABLES if k != "stats"])


@bp.route("/reports/<name>.<fmt>")
@staff
def export(name, fmt):
    if name not in TABLES or fmt not in ("xlsx", "csv"):
        abort(404)
    lang = L()
    title_key, headers, rows = TABLES[name](request.args)
    stamp = datetime.now().strftime("%Y%m%d")
    if fmt == "xlsx":
        data, mime = to_xlsx(title_key, headers, rows, lang), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        data, mime = to_csv(headers, rows, lang), "text/csv; charset=utf-8"
    return Response(data, mimetype=mime, headers={
        "Content-Disposition": f"attachment; filename={name}_{stamp}.{fmt}", "Cache-Control": "no-store"})
