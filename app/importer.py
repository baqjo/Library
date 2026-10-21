import csv
import io
import re
import xml.etree.ElementTree as ET
from datetime import date
from openpyxl import load_workbook

COLS = ["student_no", "name_ar", "name_en", "email", "stage", "class", "section", "type", "barcode"]
STAFF_WORDS = {"staff", "employee", "teacher", "موظف", "موظفة", "معلم", "معلمة", "اداري", "إداري", "إدارية"}


def is_staff_row(row):
    return (row.get("type", "") or "").strip().lower() in STAFF_WORDS


def _cell(c):
    if c is None:
        return ""
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c).strip()


def read_rows(file_storage):
    """Return list of dicts with normalized lowercase headers from CSV or XLSX."""
    name = (file_storage.filename or "").lower()
    data = file_storage.read()
    rows = []
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(it, [])]
        for r in it:
            if r is None or all(c in (None, "") for c in r):
                continue
            rows.append({header[i]: _cell(c) for i, c in enumerate(r) if i < len(header)})
    else:
        text = data.decode("utf-8-sig", errors="replace")
        rdr = csv.DictReader(io.StringIO(text))
        for r in rdr:
            rows.append({(k or "").strip().lower(): (v or "").strip() for k, v in r.items()})
    return rows


def template_csv():
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(COLS)
    w.writerow(["1001", "أحمد محمد", "Ahmed Mohammed", "ahmed@school.edu", "الابتدائية", "الصف الأول", "أ", "student", ""])
    w.writerow(["T-01", "سلمى خالد", "Salma Khaled", "salma@school.edu", "", "", "", "staff", ""])
    return "\ufeff" + out.getvalue()


_AR = re.compile(r"[\u0600-\u06FF]")


def read_destiny_report(data):
    """Parse a Follett Destiny 'Custom Report' XML export (<DestinyCustomReport><Row>...).
    Returns (rows, report_copies): one dict per title (title, author, year, edition, category, language) and the
    total copy count stated in the report. The report carries no ISBN / barcode / call number, so copies cannot be built."""
    rows, copies, seen = [], 0, set()
    root = ET.fromstring(data)
    for r in root.iter("Row"):
        f = {c.tag: (c.text or "").strip() for c in r}
        title = f.get("title", "").strip(" /:;,")
        if not title:
            continue
        try:
            copies += int(f.get("statsCurrentCopyCount") or 0)
        except ValueError:
            pass
        year = f.get("pubYear", "")
        if not (year.isdigit() and 1400 <= int(year) <= date.today().year + 1):
            year = ""
        author = f.get("author", "").strip(" /:;,")
        mat = f.get("materialType", "")
        rows.append({"title": title, "author": author, "year": year, "edition": f.get("edition", ""),
                     "category": "" if mat in ("", "Book") else mat,
                     "language": "ar" if _AR.search(title) else "en"})
    return rows, copies
